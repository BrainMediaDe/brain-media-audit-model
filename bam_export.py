"""
Brain-Media Audit Model (BAM) Core - bam_export.py

Copyright (c) 2026 Dr. Holger Reibold / Brain-Media.de
https://brain-media.de

Lizenz: GNU Affero General Public License v3.0 (AGPLv3)
Siehe LICENSE-Datei im Repository-Root.

Teil von "Brain-Media Audit Model (BAM) Core"
https://github.com/BrainMediaDe/brain-media-audit-model

Export der Gap-Analyse in Formate, die in Audit-Arbeitspapieren
weiterverwendbar sind: CSV, XLSX und JSON.

Das Modul ist zustandslos. Der Bewertungsstand wird uebergeben, nicht
gespeichert - BAM Core 2.0.x haelt den Zustand im Browser
(localStorage), eine serverseitige Persistenz kommt erst mit dem
Instanz-Layer in 2.2.

XLSX benoetigt openpyxl:
    pip3 install openpyxl --break-system-packages
Ohne openpyxl funktionieren CSV und JSON unveraendert weiter.
"""

import csv
import io
import json
from datetime import datetime, timezone

BASE_URL = "https://bam.brain-media.de"

# Antwortwerte des Dashboards (localStorage-Schluessel 'bam_gaps')
ANSWER_LABELS = {
    "ja": "erfüllt",
    "teil": "teilweise erfüllt",
    "nein": "nicht erfüllt",
}
NOT_ASSESSED = "nicht bewertet"

# Spalten des Arbeitspapiers, Reihenfolge ist bewusst gewaehlt:
# Identifikation, Anforderung, Bewertung, Feststellung, dann Planung.
COLUMNS = [
    ("bam_id", "BAM-ID"),
    ("regulation", "Regulierung"),
    ("article", "Fundstelle"),
    ("requirement", "Anforderung"),
    ("gap_question", "Prüffrage"),
    ("answer", "Bewertung"),
    ("finding", "Feststellung / Hinweis"),
    ("priority", "Priorität"),
    ("risk_score", "Risikoscore"),
    ("fine_max", "Bußgeldrahmen"),
    ("remediation", "Maßnahme"),
    ("effort", "Aufwand"),
    ("pt_min", "PT min"),
    ("pt_max", "PT max"),
    ("deadline", "Zeithorizont"),
    ("control", "Control"),
    ("control_type", "Control-Art"),
    ("evidence_type", "Nachweis"),
    ("iso27001", "ISO 27001:2022"),
    ("cross_refs", "Querverweise"),
    ("bam_uri", "Referenz"),
]


def _fine(risk):
    f = (risk or {}).get("regulatory_fine") or {}
    if not f:
        return ""
    parts = []
    mx = f.get("max_eur")
    if isinstance(mx, (int, float)):
        parts.append(f"{mx:,.0f} EUR".replace(",", "."))
    elif mx:
        parts.append(str(mx))
    if f.get("max_pct"):
        parts.append(str(f["max_pct"]))
    return " oder ".join(parts)


def build_rows(db, gaps, only_gaps=False):
    """
    Erzeugt die Zeilen des Arbeitspapiers.

    gaps: {bam_id: 'ja'|'teil'|'nein'} - der Stand aus dem Dashboard.
    only_gaps: True liefert nur Objekte mit Feststellung (teil/nein).
    """
    gaps = gaps or {}
    rows = []
    for o in db.get("objects", []):
        bid = o.get("bam_id")
        req = o.get("requirement") or {}
        gap = o.get("gap_check") or {}
        rem = o.get("remediation") or {}
        risk = o.get("risk") or {}
        ctrl = o.get("control") or {}
        evid = o.get("evidence") or {}
        iso = o.get("iso27001_mapping") or {}
        cost = rem.get("cost_estimate") or {}

        answer = gaps.get(bid)
        if only_gaps and answer not in ("teil", "nein"):
            continue

        if answer == "nein":
            finding = gap.get("if_no", "")
        elif answer == "teil":
            finding = gap.get("if_partial", "")
        elif answer == "ja":
            finding = gap.get("if_yes", "")
        else:
            finding = ""

        rows.append({
            "bam_id": bid,
            "regulation": o.get("regulation", ""),
            "article": o.get("article", ""),
            "requirement": req.get("text", ""),
            "gap_question": gap.get("question", ""),
            "answer": ANSWER_LABELS.get(answer, NOT_ASSESSED),
            "finding": finding,
            "priority": ctrl.get("priority") or req.get("priority") or "",
            "risk_score": risk.get("score", ""),
            "fine_max": _fine(risk),
            "remediation": rem.get("summary", ""),
            "effort": rem.get("effort", ""),
            "pt_min": cost.get("pt_min", ""),
            "pt_max": cost.get("pt_max", ""),
            "deadline": rem.get("deadline", ""),
            "control": ctrl.get("measure", ""),
            "control_type": ctrl.get("type", ""),
            "evidence_type": evid.get("type", ""),
            "iso27001": ", ".join(iso.get("controls", []) or []),
            "cross_refs": ", ".join(o.get("cross_refs", []) or []),
            "bam_uri": f"{BASE_URL}/id/{bid}" if bid else "",
        })
    return rows


def build_summary(db, gaps):
    """
    Kennzahlen je Regulierung und gesamt. Die Punktlogik entspricht
    der des Dashboards: erfüllt = 1, teilweise = 0,5, nicht erfüllt = 0,
    bezogen auf alle Objekte mit Prüffrage.
    """
    gaps = gaps or {}
    per_reg = {}
    for o in db.get("objects", []):
        if not (o.get("gap_check") or {}).get("question"):
            continue
        reg = o.get("regulation", "Sonstige")
        a = gaps.get(o.get("bam_id"))
        e = per_reg.setdefault(reg, {
            "regulation": reg, "total": 0, "bewertet": 0,
            "erfuellt": 0, "teilweise": 0, "nicht_erfuellt": 0, "punkte": 0.0,
        })
        e["total"] += 1
        if a:
            e["bewertet"] += 1
        if a == "ja":
            e["erfuellt"] += 1
            e["punkte"] += 1
        elif a == "teil":
            e["teilweise"] += 1
            e["punkte"] += 0.5
        elif a == "nein":
            e["nicht_erfuellt"] += 1

    rows = []
    for e in sorted(per_reg.values(), key=lambda x: x["regulation"]):
        e["score_pct"] = round(e["punkte"] / e["total"] * 100) if e["total"] else 0
        rows.append(e)

    tot = {
        "regulation": "Gesamt",
        "total": sum(e["total"] for e in rows),
        "bewertet": sum(e["bewertet"] for e in rows),
        "erfuellt": sum(e["erfuellt"] for e in rows),
        "teilweise": sum(e["teilweise"] for e in rows),
        "nicht_erfuellt": sum(e["nicht_erfuellt"] for e in rows),
        "punkte": sum(e["punkte"] for e in rows),
    }
    tot["score_pct"] = round(tot["punkte"] / tot["total"] * 100) if tot["total"] else 0
    return rows, tot


def build_meta(db, organisation="", assessor="", note=""):
    return {
        "exportiert_am": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
        "modellversion": str(db.get("version") or db.get("schema_version") or ""),
        "modellstand": str(db.get("released") or ""),
        "schema_version": str(db.get("schema_version") or ""),
        "organisation": organisation or "",
        "bewertet_durch": assessor or "",
        "anmerkung": note or "",
        "quelle": db.get("_attribution", "Brain-Media Audit Model (BAM)"),
        "lizenz": db.get("_license", ""),
    }


# ---------------------------------------------------------------- CSV

def to_csv(rows, delimiter=";"):
    """
    CSV mit BOM und Semikolon, damit Excel unter Windows die Datei
    ohne Importdialog korrekt oeffnet.
    """
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL,
                   lineterminator="\r\n")
    w.writerow([label for _, label in COLUMNS])
    for r in rows:
        w.writerow([r.get(key, "") for key, _ in COLUMNS])
    return buf.getvalue().encode("utf-8-sig")


# --------------------------------------------------------------- JSON

def to_json(rows, meta, summary_rows, summary_total):
    payload = {
        "meta": meta,
        "zusammenfassung": {"je_regulierung": summary_rows, "gesamt": summary_total},
        "positionen": rows,
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


# --------------------------------------------------------------- XLSX

def to_xlsx(rows, meta, summary_rows, summary_total):
    """
    Arbeitsmappe mit drei Blättern: Deckblatt, Zusammenfassung,
    Gap-Analyse. Benoetigt openpyxl.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    INK = "FF0F0F0D"
    HEAD_FILL = PatternFill("solid", fgColor="FFF2F0EC")
    RULE = Side(style="thin", color="FFE2DED6")
    BORDER = Border(bottom=RULE)
    H = Font(name="Calibri", size=10, bold=True, color=INK)
    B = Font(name="Calibri", size=10, color=INK)
    TITLE = Font(name="Calibri", size=14, bold=True, color=INK)

    wb = Workbook()

    # --- Deckblatt ---
    ws = wb.active
    ws.title = "Deckblatt"
    ws["A1"] = "Gap-Analyse"
    ws["A1"].font = TITLE
    ws["A2"] = "Brain-Media Audit Model (BAM) Core"
    ws["A2"].font = B
    labels = [
        ("Organisation", meta.get("organisation")),
        ("Bewertet durch", meta.get("bewertet_durch")),
        ("Exportiert am", meta.get("exportiert_am")),
        ("Modellversion", meta.get("modellversion")),
        ("Modellstand", meta.get("modellstand")),
        ("Schemaversion", meta.get("schema_version")),
        ("Anmerkung", meta.get("anmerkung")),
        ("Quelle", meta.get("quelle")),
        ("Lizenz", meta.get("lizenz")),
        ("Bezeichner", f"Stabil und auflösbar unter {BASE_URL}/id/<BAM-ID>"),
    ]
    r = 4
    for label, value in labels:
        ws.cell(row=r, column=1, value=label).font = H
        c = ws.cell(row=r, column=2, value=value or "")
        c.font = B
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 72

    # --- Zusammenfassung ---
    ws = wb.create_sheet("Zusammenfassung")
    head = ["Regulierung", "Prüffragen", "bewertet", "erfüllt",
            "teilweise", "nicht erfüllt", "Erfüllungsgrad"]
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = H
        c.fill = HEAD_FILL
        c.border = BORDER
    r = 2
    for e in list(summary_rows) + [summary_total]:
        vals = [e["regulation"], e["total"], e["bewertet"], e["erfuellt"],
                e["teilweise"], e["nicht_erfuellt"], e["score_pct"] / 100]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name="Calibri", size=10, color=INK,
                          bold=(e is summary_total))
            if i == 7:
                c.number_format = "0 %"
        r += 1
    for i, w in enumerate([22, 12, 12, 12, 12, 14, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # --- Gap-Analyse ---
    ws = wb.create_sheet("Gap-Analyse")
    for i, (_, label) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = H
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(vertical="top")
    for ri, row in enumerate(rows, start=2):
        for ci, (key, _) in enumerate(COLUMNS, start=1):
            v = row.get(key, "")
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = B
            c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {"bam_id": 30, "regulation": 12, "article": 18, "requirement": 60,
              "gap_question": 55, "answer": 18, "finding": 55, "priority": 12,
              "risk_score": 11, "fine_max": 20, "remediation": 55, "effort": 14,
              "pt_min": 8, "pt_max": 8, "deadline": 18, "control": 45,
              "control_type": 16, "evidence_type": 28, "iso27001": 24,
              "cross_refs": 34, "bam_uri": 50}
    for i, (key, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(key, 18)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 1)}"
    )

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def filename(fmt, organisation="", only_gaps=False):
    stamp = datetime.now().strftime("%Y-%m-%d")
    org = "".join(ch if ch.isalnum() or ch in "-_" else "-"
                  for ch in (organisation or "")).strip("-")
    parts = ["bam", "gap-analyse" if not only_gaps else "gap-feststellungen"]
    if org:
        parts.append(org)
    parts.append(stamp)
    return "-".join(parts) + "." + fmt
